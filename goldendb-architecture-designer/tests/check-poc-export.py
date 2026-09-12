"""Compare the actual POC workbook with its exported sheet model."""
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

root = Path(__file__).resolve().parents[2] / 'outputs/goldendb-remediation-20260906/poc-density/evidence'
expected = json.loads((root / 'poc-sheets.json').read_text())
workbook = openpyxl.load_workbook(root / 'poc.xlsx', data_only=False)
differences = []
checked = 0
assert workbook.sheetnames == [sheet['name'] for sheet in expected]
with zipfile.ZipFile(root / 'poc.xlsx') as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith(('.xml', '.rels')):
            ET.fromstring(archive.read(name))
for sheet in expected:
    ws = workbook[sheet['name']]
    for row_index, row in enumerate(sheet['rows'], 1):
        for col_index, value in enumerate(row, 1):
            cell = ws.cell(row_index, col_index)
            checked += 1
            if cell.data_type in ('e', 'f') or (cell.value or '') != (value or ''):
                differences.append([ws.title, cell.coordinate, value, cell.value])
result = {'checkedCells': checked, 'differences': differences, 'nativeOfficeChecked': False}
(root / 'poc-workbook-validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps(result, ensure_ascii=False))
assert not differences
