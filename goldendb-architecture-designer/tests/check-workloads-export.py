"""Validate actual workload export against the application workbook model."""
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

root = Path(__file__).resolve().parents[2] / 'outputs/goldendb-remediation-20260906/workloads/evidence'
expected = json.loads((root / 'sheets.json').read_text())
workbook = openpyxl.load_workbook(root / 'workloads.xlsx')
assert workbook.sheetnames == [sheet['name'] for sheet in expected]
with zipfile.ZipFile(root / 'workloads.xlsx') as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith(('.xml', '.rels')):
            ET.fromstring(archive.read(name))
checked = 0
for sheet in expected:
    ws = workbook[sheet['name']]
    for row_index, row in enumerate(sheet['rows'], 1):
        for col_index, value in enumerate(row, 1):
            cell = ws.cell(row_index, col_index)
            assert cell.data_type not in ('e', 'f')
            assert (cell.value or '') == (value or ''), (ws.title, cell.coordinate)
            checked += 1
result = {'checkedCells': checked, 'differences': 0, 'nativeOfficeChecked': False}
(root / 'workbook-validation.json').write_text(json.dumps(result, indent=2))
print(json.dumps(result))
